import csv
import logging
import os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

logging.basicConfig(level=logging.INFO, format="[EXCEL] %(message)s")
log = logging.getLogger(__name__)

REPORTS_DIR = "/opt/airflow/reports"

HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


def _csv_path(name: str) -> str:
    return os.path.join(REPORTS_DIR, name)


def _read_csv(filename: str) -> tuple[list[str], list[list[str]]]:
    path = _csv_path(filename)
    if not os.path.exists(path):
        log.warning("  ! %s no encontrado, saltando", filename)
        return [], []
    with open(path, newline="", encoding="utf-8") as f:
        reader = csv.reader(f)
        rows = list(reader)
    if not rows:
        return [], []
    return rows[0], rows[1:]


def _add_sheet(wb: Workbook, title: str, headers: list[str], rows: list[list[str]]):
    if not headers:
        return
    ws = wb.create_sheet(title=title[:31])
    ws.append(headers)
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
        cell.border = THIN_BORDER
    for row in rows:
        ws.append(row)
    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=0)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 3, 40)
    log.info("  %s: %d filas", title, len(rows))


SHEETS = [
    ("Tendencias - Estados", "tendencias_estados.csv"),
    ("Tendencias - Tamaño Grupo", "tendencias_avg_party.csv"),
    ("Tendencias - Categorías", "tendencias_categorias.csv"),
    ("Tendencias - Ranking", "tendencias_ranking.csv"),
    ("Horarios - Por Hora", "horarios_por_hora.csv"),
    ("Horarios - Por Día", "horarios_por_dia.csv"),
    ("Horarios - Ventanas Pico", "horarios_ventanas_pico.csv"),
    ("Crecimiento Mensual", "crecimiento_mensual.csv"),
]


def main():
    print("=" * 60)
    print("EXPORT — Generando reportes Excel desde CSV")
    print("=" * 60)

    wb = Workbook()
    wb.remove(wb.active)

    for sheet_name, csv_file in SHEETS:
        headers, rows = _read_csv(csv_file)
        _add_sheet(wb, sheet_name, headers, rows)

    if len(wb.sheetnames) == 0:
        log.warning("No se encontraron CSVs para convertir.")
        return

    output_path = _csv_path("reportes_analitica.xlsx")
    wb.save(output_path)
    print(f"\n  Excel generado: {output_path}")
    print(f"  Hojas: {', '.join(wb.sheetnames)}")
    print("=" * 60)


if __name__ == "__main__":
    main()
